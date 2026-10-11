"""L7: Q179 '410 million total genomic differences, not 35 million' -- which accounting of Yoo 2025's supplementary
sheets gives 410 M (and Day's earlier 14.9%), using the three sheets Day now names ('pages 14, 19, and 24').
Low-stakes tier (re-confirmation / bookkeeping; one combined review). Runs locally (one process, < 1 min).
Pre-registered: committed BEFORE it is run. DISCLOSURE: the values of sheets 14 and 19 (and the column layout of
sheet 24) were looked at during the derivation step, before this file was written; the predictions below are
therefore partly informed (reading R1 was computed by hand: 13.94% x 2,900.6 Mb = 404 Mb).

Run: research/.venv/bin/python -I research/checks/l7_410m_accounting.py [MOESM4.xlsx]
Input: Yoo et al. 2025 Supplementary Tables (MOESM4.xlsx, local sources/raw/sources/supp/Yoo2025/, gitignored).
Day's words (post 2026-10-11, quotes-day.md Q179 and body ¶ on the supplement): '410 million total genomic
differences, not 35 million'; 'SNV divergence and structural divergence are reported separately, in different sheets
of an 82-sheet supplementary spreadsheet, computed by three different tools on three different bases ... buried on
pages 14, 19, and 24 ... they aren't additive and no total is either calculated or provided.' Earlier (A3a, 2nd
edition): '14.9 percent, with 410 million base pairs'; '35 million SNVs, 1,140 interspecific inversions, and
approximately 187 megabases of structurally divergent regions, for a total of approximately 410 million'.
Sheets: 14 = Table III.14 'SNV-GAP divergence' (wfmash, 1 Mb segments, mean/median); 19 = Table III.19 'Alignment
coverage of T2T-CHM13 (hs1) in the 8-way primary Progressive Cactus alignment'; 24 = Table V.24 'Ape structurally
divergent regions' (per-region rows; SDR_Size by Lineage and Haplotype).

READINGS COMPUTED (each in bp and as % of the hs1 length used)
  R1 (mean SNV div + mean gap div) x length, sheet 14 human-chimp row; autosomes; A+X; A+X+Y.
  R2 (1 - identical) x length, sheet 19 Chimp rows: 'Identical', 'Identical 1:1'; A, A+X, A+X+Y.
  R3 35 M + 2 x SDR(human lineage) and 35 M + SDR(human) + SDR(chimp), sheet 24, per haplotype and averaged.
  R4 mean SNV x length + SDR(human) + SDR(chimp).
  R5 base rate: all 1-3 element sums from the component pool within +/-1% of 410 M, against the same count for
     200 random targets uniform in 300-520 M (as in R4 GAP-07b).
PREDICTIONS (fixed now)
  Day side: a reading built only from sheets 14/19/24 gives 410 M +/- 2% (credit: 410 M is reconstructible from the
    SI, which would move A3x1's 'not-found' toward a sourced derivation).
  Critic side: any reproducing reading sums base pairs from overlapping measures across both genomes (gap bp and SDR
    bp overlap; R1/R2 count unaligned or non-identical bases, not mutational events); the event count (R4 GAP-07t,
    48.3 M both lineages) stays >= 8x lower; '35 million' is a CSAC SNV count, so '410 M, not 35 M' compares bp with
    events.
  Audit: R1 autosomal = 404 M (13.94%), within 2% of 410 M; no reading gives both 410 M (+/-2%) and 14.9% (+/-0.3
    points) on an hs1 denominator; R5 hits at 410 M are within 2x of the random-target rate (a 410 M match alone is
    weak evidence of the route).
"""
import hashlib
import itertools
import random
import sys

import openpyxl

PATH = sys.argv[1] if len(sys.argv) > 1 else "sources/raw/sources/supp/Yoo2025/MOESM4.xlsx"
TARGET = 410e6


def rows(ws):
    return [r for r in ws.iter_rows(values_only=True)]


def main():
    print("MOESM4 sha256", hashlib.sha256(open(PATH, "rb").read()).hexdigest())
    wb = openpyxl.load_workbook(PATH, read_only=True)
    t14 = rows(wb["14"])
    hc = [r for r in t14 if r[0] == "mPanTro3#1/2" and r[1] == "hg002#M/P"][0]
    snv = dict(A=hc[2], X=hc[3], Y=hc[4])
    gap = dict(A=hc[5], X=hc[6], Y=hc[7])
    print("sheet 14 human-chimp mean SNV", snv, "mean gap", gap)
    t19 = rows(wb["19"])
    L = {}
    ident = {}
    for r in t19[1:]:
        if r[2] in ("Chimp", "Chimp Pri"):
            reg = {"Autosomes": "A", "X": "X", "Y": "Y"}[r[0]]
            L[reg] = r[1]
            ident[reg] = dict(aligned=r[3], identical=r[4], aligned11=r[5], identical11=r[6])
    print("sheet 19 hs1 lengths", L, "chimp rows", ident)
    sdr = {}
    for r in wb["24"].iter_rows(min_row=10, values_only=True):
        if r[0] is None or not isinstance(r[6], (int, float)):
            continue
        sdr[(r[4], r[3])] = sdr.get((r[4], r[3]), 0) + r[6]
    print("sheet 24 SDR bp by (lineage, haplotype):", {k: round(v / 1e6, 1) for k, v in sorted(sdr.items())})
    out = []
    regs = [("A",), ("A", "X"), ("A", "X", "Y")]
    for rg in regs:
        Ln = sum(L[x] for x in rg)
        v = sum((snv[x] + gap[x]) * L[x] for x in rg)
        out.append((f"R1 (SNV+gap) x L [{'+'.join(rg)}]", v, Ln))
        for k in ("identical", "identical11", "aligned", "aligned11"):
            v2 = sum((1 - ident[x][k] / 100) * L[x] for x in rg)
            out.append((f"R2 (1-{k}) x L [{'+'.join(rg)}]", v2, Ln))
        out.append((f"R4a SNV x L [{'+'.join(rg)}]", sum(snv[x] * L[x] for x in rg), Ln))
    lineages = sorted({k[0] for k in sdr})
    hum = [k for k in sdr if k[0] in ("HSA", "Homo", "Human", "human")]
    chi = [k for k in sdr if k[0] in ("PTR", "Pan troglodytes", "Chimpanzee", "chimp")]
    print("lineages in sheet 24:", lineages, "human keys", hum, "chimp keys", chi)
    Ltot = sum(L.values())
    if hum:
        h_avg = sum(sdr[k] for k in hum) / len(hum)
        for k in hum:
            out.append((f"R3 35M + 2 x SDR{k}", 35e6 + 2 * sdr[k], Ltot))
        out.append(("R3 35M + 2 x SDR(human mean of haplotypes)", 35e6 + 2 * h_avg, Ltot))
        for kh, kc in itertools.product(hum, chi):
            out.append((f"R3 35M + SDR{kh} + SDR{kc}", 35e6 + sdr[kh] + sdr[kc], Ltot))
            out.append((f"R4 SNVxL[A+X+Y] + SDR{kh} + SDR{kc}", sum(snv[x] * L[x] for x in L) + sdr[kh] + sdr[kc], Ltot))
    print("\nreading | bp (M) | % of hs1 length used | vs 410 M")
    for name, v, Ln in out:
        flag = " <== within 2% of 410 M" if abs(v / TARGET - 1) <= 0.02 else ""
        flag += " ; 14.9% +/- 0.3" if abs(100 * v / Ln - 14.9) <= 0.3 else ""
        print(f"{name:55s} {v / 1e6:8.1f} {100 * v / Ln:6.2f}% {v / TARGET:6.3f}{flag}")
    print("14.9% denominators: 410 M / 0.149 = %.0f Mb; hs1 A = %.0f, A+X = %.0f, A+X+Y = %.0f Mb" % (
        TARGET / 0.149 / 1e6, L['A'] / 1e6, (L['A'] + L['X']) / 1e6, Ltot / 1e6))
    pool = [v for _, v, _ in out] + [35e6, 1140, 48.27e6, 42.87e6, 5.40e6] + list(sdr.values())
    pool = sorted(set(round(p) for p in pool))

    def hits(t):
        n = 0
        for r in (1, 2, 3):
            for c in itertools.combinations(pool, r):
                if abs(sum(c) / t - 1) <= 0.01:
                    n += 1
        return n
    h410 = hits(TARGET)
    random.seed(20261011)
    rt = [hits(random.uniform(300e6, 520e6)) for _ in range(200)]
    print(f"\nR5 base rate: pool {len(pool)} values; 1-3 element sums within 1% of 410 M: {h410}; random targets "
          f"300-520 M (200 draws): mean {sum(rt) / len(rt):.1f}, max {max(rt)}; ratio {h410 / max(1e-9, sum(rt) / len(rt)):.2f}")
    print("Ratios: 410/35 = %.2f; 410/48.27 (GAP-07t events, both lineages) = %.2f; 14.9/1.23 = %.2f" % (
        410 / 35, 410 / 48.27, 14.9 / 1.23))


if __name__ == "__main__":
    main()
